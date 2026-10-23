import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # Color Palette: Dark Black & Light Blue
    COLOR_BG = RGBColor(10, 13, 20)           # #0A0D14 (Deep dark black)
    COLOR_CARD = RGBColor(20, 30, 50)         # #141E32 (Dark Navy card)
    COLOR_CARD_BORDER = RGBColor(30, 58, 95)  # #1E3A5F
    COLOR_LIGHT_BLUE = RGBColor(56, 189, 248) # #38BDF8 (Vibrant light blue / cyan)
    COLOR_ACCENT_BLUE = RGBColor(96, 165, 250)# #60A5FA (Soft sky blue)
    COLOR_MUTED_BLUE = RGBColor(147, 197, 253)# #93C5FD (Light pastel blue)
    COLOR_TEXT_WHITE = RGBColor(248, 250, 252)# #F8FAFC
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8
    
    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.color.rgb = COLOR_BG
        return bg_shape

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Tag badge
        tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.4))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_LIGHT_BLUE
        
        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_WHITE
        
        # Subtitle
        if subtitle_text:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.4))
            p_sub = tb_sub.text_frame.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(slide, left, top, width, height, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # =========================================================================
    # SLIDE 1: Title & Team Details
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    # Project Badge
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.3), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "SMART MUNICIPAL INFRASTRUCTURE & AI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_LIGHT_BLUE
    
    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.2))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = "CivicFix"
    p_title.font.size = Pt(56)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE
    
    # Subtitle / Tagline
    tb_tagline = s1.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(11.3), Inches(0.6))
    p_tagline = tb_tagline.text_frame.paragraphs[0]
    p_tagline.text = "Scan. Report. Fixed. — AI-Verified Civic Redressal & GIS Dispatch Platform"
    p_tagline.font.size = Pt(18)
    p_tagline.font.color.rgb = COLOR_ACCENT_BLUE

    # Divider Line
    line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.6), Inches(11.3), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.color.rgb = COLOR_CARD_BORDER

    # Team Card
    add_card(s1, Inches(1.0), Inches(3.9), Inches(11.3), Inches(2.8))
    
    tb_team_header = s1.shapes.add_textbox(Inches(1.3), Inches(4.1), Inches(10.7), Inches(0.5))
    p_th = tb_team_header.text_frame.paragraphs[0]
    p_th.text = "TEAM: DevOps Titans"
    p_th.font.size = Pt(20)
    p_th.font.bold = True
    p_th.font.color.rgb = COLOR_LIGHT_BLUE

    # Team Members Grid (4 Columns)
    members = [
        ("Kevin", "Team Lead / Cloud & DevOps", "AWS Architecture & Deployment"),
        ("Member 2", "Backend & AI Engineer", "Spring Boot & Gemini Vision AI"),
        ("Member 3", "Frontend Architect", "React 19 & Interactive GIS Maps"),
        ("Member 4", "Database & QA Engineer", "MySQL Data Integrity & Testing")
    ]
    
    col_w = Inches(2.55)
    gap = Inches(0.2)
    start_x = Inches(1.3)
    
    for i, (name, role, task) in enumerate(members):
        mx = start_x + i * (col_w + gap)
        tb_m = s1.shapes.add_textbox(mx, Inches(4.7), col_w, Inches(1.8))
        tf = tb_m.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = name
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_WHITE
        
        p2 = tf.add_paragraph()
        p2.text = role
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_ACCENT_BLUE
        
        p3 = tf.add_paragraph()
        p3.text = task
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 2: Problem Statement & Solution Provided
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Challenge & Innovation", "Problem Statement & Solution", "Bridging the gap between citizen grievances and municipal authorities")
    
    card_w = Inches(5.6)
    card_h = Inches(4.8)
    top_y = Inches(2.0)
    
    # Left Card: Problem Statement
    add_card(s2, Inches(0.8), top_y, card_w, card_h)
    tb_prob = s2.shapes.add_textbox(Inches(1.1), top_y + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5))
    tf_p = tb_prob.text_frame
    tf_p.word_wrap = True
    
    hp = tf_p.paragraphs[0]
    hp.text = "THE PROBLEM STATEMENT"
    hp.font.size = Pt(18)
    hp.font.bold = True
    hp.font.color.rgb = RGBColor(248, 113, 113) # Coral red accent
    
    problems = [
        ("Slow & Fragmented Redressal", "Citizens struggle with bureaucratic, outdated complaint portals with zero visibility into resolution status."),
        ("Spam, Fakes & Overloaded Staff", "Municipal bodies waste hundreds of hours filtering non-civic photos, unrelated images, and duplicate complaints."),
        ("Vague & Inaccurate Locations", "Complaints lack pinpoint coordinates, causing municipal field crews to wander looking for potholes and dumps."),
        ("Citizen Privacy Violations", "Unregulated public photos often capture pedestrians' faces, creating legal and privacy liabilities.")
    ]
    
    for title, desc in problems:
        pt = tf_p.add_paragraph()
        pt.text = f"• {title}:"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_WHITE
        
        pd = tf_p.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(11)
        pd.font.color.rgb = COLOR_TEXT_MUTED

    # Right Card: Solution Provided
    add_card(s2, Inches(6.9), top_y, card_w, card_h)
    tb_sol = s2.shapes.add_textbox(Inches(7.2), top_y + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5))
    tf_s = tb_sol.text_frame
    tf_s.word_wrap = True
    
    hs = tf_s.paragraphs[0]
    hs.text = "THE CIVICFIX SOLUTION"
    hs.font.size = Pt(18)
    hs.font.bold = True
    hs.font.color.rgb = COLOR_LIGHT_BLUE
    
    solutions = [
        ("AI Vision Verification (Gemini 1.5)", "Instantly validates genuine road, garbage, and infrastructure issues, generating auto-descriptions."),
        ("Automated Privacy Shield", "AI scans uploaded photos; immediately rejects any image containing human faces to protect citizen privacy."),
        ("Hyper-Local Locality Intelligence", "Pre-mapped GPS tagging for Eachanari, Gandhipuram, etc., with interactive maps and route navigation."),
        ("Unified Department Hub", "Department-segregated review queue (Roads, Sanitation, Water) with live status tracking for citizens.")
    ]
    
    for title, desc in solutions:
        st = tf_s.add_paragraph()
        st.text = f"✔ {title}:"
        st.font.size = Pt(12)
        st.font.bold = True
        st.font.color.rgb = COLOR_ACCENT_BLUE
        
        sd = tf_s.add_paragraph()
        sd.text = f"   {desc}"
        sd.font.size = Pt(11)
        sd.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: AWS Services Used
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Cloud Infrastructure", "AWS Cloud Architecture", "Scalable, resilient, and enterprise-grade deployment on Amazon Web Services")

    aws_services = [
        ("Amazon EC2", "Compute & Hosting", [
            "Hosts Dockerized multi-tier stack (Nginx, Spring Boot, MySQL)",
            "Ubuntu 24.04 LTS instance running in public subnet",
            "High availability with container auto-restart policies"
        ]),
        ("Amazon S3", "Object Cloud Storage", [
            "Durable cloud bucket for high-res civic issue photos",
            "Automatic fallback to local container volume on network fail",
            "Direct secure CDN URLs for citizen & admin inspection"
        ]),
        ("AWS VPC & Security", "Network Isolation", [
            "Strict Security Group firewalls (Port 80 HTTP, Port 8080 API)",
            "Container-internal bridge network for MySQL isolation",
            "Safe traffic ingress protecting database credentials"
        ]),
        ("AWS IAM", "Access Management", [
            "Least-privilege security roles for EC2-to-S3 uploads",
            "Credential-less cloud authorization via instance profiles",
            "Encrypted environment variable secrets management"
        ])
    ]

    card_w3 = Inches(2.7)
    card_h3 = Inches(4.7)
    start_x3 = Inches(0.8)
    gap3 = Inches(0.3)
    top_y3 = Inches(2.0)

    for i, (service, role, points) in enumerate(aws_services):
        cx = start_x3 + i * (card_w3 + gap3)
        add_card(s3, cx, top_y3, card_w3, card_h3)
        
        tb = s3.shapes.add_textbox(cx + Inches(0.2), top_y3 + Inches(0.25), card_w3 - Inches(0.4), card_h3 - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        h = tf.paragraphs[0]
        h.text = service
        h.font.size = Pt(17)
        h.font.bold = True
        h.font.color.rgb = COLOR_LIGHT_BLUE
        
        r = tf.add_paragraph()
        r.text = role
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_ACCENT_BLUE
        
        tf.add_paragraph().text = "" # spacer
        
        for p in points:
            pt = tf.add_paragraph()
            pt.text = f"• {p}"
            pt.font.size = Pt(10.5)
            pt.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Benefits & Feasibilities
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Impact & Viability", "Key Benefits & Feasibility Analysis", "Demonstrated municipal ROI, operational savings, and effortless scalability")

    # Left: Key Benefits
    add_card(s4, Inches(0.8), top_y, card_w, card_h)
    tb_b = s4.shapes.add_textbox(Inches(1.1), top_y + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    
    hb = tf_b.paragraphs[0]
    hb.text = "CORE BENEFITS & IMPACT"
    hb.font.size = Pt(18)
    hb.font.bold = True
    hb.font.color.rgb = COLOR_LIGHT_BLUE
    
    benefits = [
        ("80% Reduction in Triage Time", "Automated AI classification directs complaints straight to Roads, Sanitation, or Water officials without manual triage."),
        ("Elimination of False Reports", "Computer vision model rejects fake photos, non-civic uploads, and duplicate entries at the submission gate."),
        ("Turn-by-Turn Field Dispatch", "Integrated GPS coordinates and Google Maps directions guide municipal maintenance teams directly to the spot."),
        ("Transparent Citizen Tracking", "Real-time status updates (Not Reviewed -> Reviewed -> Resolved) build public trust in local governance.")
    ]
    for b_title, b_desc in benefits:
        p1 = tf_b.add_paragraph()
        p1.text = f"★ {b_title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_ACCENT_BLUE
        p2 = tf_b.add_paragraph()
        p2.text = f"   {b_desc}"
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Feasibility Analysis
    add_card(s4, Inches(6.9), top_y, card_w, card_h)
    tb_f = s4.shapes.add_textbox(Inches(7.2), top_y + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    
    hf = tf_f.paragraphs[0]
    hf.text = "PROJECT FEASIBILITY"
    hf.font.size = Pt(18)
    hf.font.bold = True
    hf.font.color.rgb = COLOR_LIGHT_BLUE
    
    feasibilities = [
        ("Technical Feasibility: HIGH", "Built on robust enterprise technologies (Spring Boot, React 19, Docker, MySQL) with zero proprietary vendor lock-in."),
        ("Financial Viability: LOW COST", "Runs efficiently on a single EC2 micro/small instance with pay-as-you-go S3 storage, minimizing public expenditure."),
        ("Operational Simplicity", "Zero installation needed for citizens or staff — operates entirely via browser across mobiles, tablets, and desktops."),
        ("Scalability & Extensibility", "Modular REST architecture ready to connect to smart city IoT sensors, WhatsApp grievance bots, and state-wide portals.")
    ]
    for f_title, f_desc in feasibilities:
        p1 = tf_f.add_paragraph()
        p1.text = f"✔ {f_title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_WHITE
        p2 = tf_f.add_paragraph()
        p2.text = f"   {f_desc}"
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 5: Thank You
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)

    # Center Card
    center_w = Inches(10.0)
    center_h = Inches(5.2)
    center_x = (prs.slide_width - center_w) / 2
    center_y = (prs.slide_height - center_h) / 2
    add_card(s5, center_x, center_y, center_w, center_h)

    tb_thank = s5.shapes.add_textbox(center_x + Inches(0.5), center_y + Inches(0.6), center_w - Inches(1.0), Inches(4.0))
    tf_t = tb_thank.text_frame
    tf_t.word_wrap = True

    p_thx = tf_t.paragraphs[0]
    p_thx.text = "Thank You!"
    p_thx.font.size = Pt(50)
    p_thx.font.bold = True
    p_thx.font.color.rgb = COLOR_LIGHT_BLUE
    p_thx.alignment = PP_ALIGN.CENTER

    p_subthx = tf_t.add_paragraph()
    p_subthx.text = "Building Cleaner, Safer & Smarter Neighborhoods Together"
    p_subthx.font.size = Pt(20)
    p_subthx.font.color.rgb = COLOR_TEXT_WHITE
    p_subthx.alignment = PP_ALIGN.CENTER

    tf_t.add_paragraph().text = "" # spacer

    p_team = tf_t.add_paragraph()
    p_team.text = "Presented by DevOps Titans"
    p_team.font.size = Pt(18)
    p_team.font.bold = True
    p_team.font.color.rgb = COLOR_ACCENT_BLUE
    p_team.alignment = PP_ALIGN.CENTER

    tf_t.add_paragraph().text = "" # spacer

    p_links = tf_t.add_paragraph()
    p_links.text = "🌐 Live Deployment: http://3.88.100.25   |   📁 Repository: github.com/Kevin-13579/civic"
    p_links.font.size = Pt(14)
    p_links.font.color.rgb = COLOR_MUTED_BLUE
    p_links.alignment = PP_ALIGN.CENTER

    p_qa = tf_t.add_paragraph()
    p_qa.text = "Questions & Feedback Welcome"
    p_qa.font.size = Pt(13)
    p_qa.font.color.rgb = COLOR_TEXT_MUTED
    p_qa.alignment = PP_ALIGN.CENTER

    output_path = "CivicFix_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
